from openpyxl import Workbook
from openpyxl.styles import Font, PatternFill, Alignment, Border, Side
from openpyxl.utils import get_column_letter

wb = Workbook()

def fill(hex_): return PatternFill("solid", start_color=hex_.lstrip("#"), end_color=hex_.lstrip("#"))
def font(bold=False, color="333333", size=11): return Font(bold=bold, color=color, size=size)
def center(): return Alignment(horizontal="center", vertical="center", wrap_text=True)
def left(): return Alignment(horizontal="left", vertical="center", wrap_text=True)
def border():
    s = Side(style="thin", color="CCCCCC")
    return Border(left=s, right=s, top=s, bottom=s)

def header_row(ws, row, cols, bg="1E3A5F", fg="FFFFFF"):
    for col, val in enumerate(cols, 1):
        c = ws.cell(row=row, column=col, value=val)
        c.fill = fill(bg); c.font = font(bold=True, color=fg); c.alignment = center(); c.border = border()

def cell(ws, row, col, val, bg=None, bold=False, color="333333", align=None):
    c = ws.cell(row=row, column=col, value=val)
    if bg: c.fill = fill(bg)
    c.font = font(bold=bold, color=color)
    c.alignment = align or left()
    c.border = border()
    return c

# ── Sheet 1: STLC Overview ────────────────────────────────────────────────────
ws1 = wb.active
ws1.title = "STLC_Overview"
ws1.sheet_view.showGridLines = False
ws1.freeze_panes = "A3"

ws1.merge_cells("A1:H1")
c = ws1["A1"]
c.value = "SOFTWARE TESTING LIFE CYCLE (STLC) — Task Manager v2.0"
c.font = font(bold=True, color="FFFFFF", size=14)
c.fill = fill("0F172A"); c.alignment = center()
ws1.row_dimensions[1].height = 38

headers = ["Phase","Objective","Entry Criteria","Exit Criteria","Activities","Owner","Duration","Status"]
header_row(ws1, 2, headers)
ws1.row_dimensions[2].height = 32

PHASE_COLORS = {
    "1. Requirement Analysis":"DBEAFE",
    "2. Test Planning":"EDE9FE",
    "3. Test Case Design":"FEF3C7",
    "4. Test Environment Setup":"D1FAE5",
    "5. Test Execution":"FCE7F3",
    "6. Test Closure":"F0F9FF",
}
STATUS_COLOR = {"Completed":"D1FAE5","In Progress":"FEF3C7","Pending":"F8FAFC"}
STATUS_TEXT  = {"Completed":"065F46","In Progress":"92400E","Pending":"94A3B8"}

phases = [
    ("1. Requirement Analysis",
     "Understand what needs to be tested",
     "BRD/SRS documents available; Stakeholders accessible",
     "All requirements reviewed & test basis signed off",
     "Review requirements\nIdentify testable features\nClarify ambiguities\nPrioritize test areas",
     "BA + QA Lead","2 days","Completed"),
    ("2. Test Planning",
     "Define scope, approach, resources & schedule",
     "Requirements analysis complete; Resources allocated",
     "Test plan document approved by stakeholders",
     "Define test scope (in/out)\nSelect test types\nChoose tools (Postman, k6)\nEstimate effort\nIdentify risks",
     "QA Lead","1 day","Completed"),
    ("3. Test Case Design",
     "Create detailed test cases & RTM",
     "Test plan approved; Requirements finalized",
     "All test cases written, reviewed & added to RTM",
     "Write test cases (TC-001–TC-015)\nCreate RTM\nPeer review\nGet sign-off",
     "QA Engineers","3 days","Completed"),
    ("4. Test Environment Setup",
     "Prepare backend, frontend & test data",
     "Test cases ready; Infrastructure available",
     "Environment verified & smoke-tested",
     "Install Python 3.11 + FastAPI\nSetup React+Vite frontend\nSeed demo data\nVerify CORS & auth\nRun smoke tests",
     "DevOps + Dev","1 day","Completed"),
    ("5. Test Execution",
     "Execute all test cases & log defects",
     "Environment ready; Test cases approved",
     "All TCs executed; All Critical/High defects fixed",
     "Execute TC-001–TC-015\nLog defects (DEF-001–DEF-003)\nRetest fixes\nUpdate RTM Pass/Fail\nRegression testing",
     "QA Engineers","3 days","In Progress"),
    ("6. Test Closure",
     "Evaluate completion & document metrics",
     "All planned tests executed; No P1 open defects",
     "Test closure report signed off; Metrics archived",
     "Compute pass rate\nWrite closure report\nLesson-learned meeting\nArchive RTM & STLC\nSign-off for release",
     "QA Lead","1 day","Pending"),
]

for i, row in enumerate(phases, 3):
    bg = PHASE_COLORS.get(row[0], "F8FAFC")
    sbg = STATUS_COLOR.get(row[7], "F8FAFC")
    stxt = STATUS_TEXT.get(row[7], "333333")
    for j, val in enumerate(row, 1):
        if j == 8:
            c2 = cell(ws1, i, j, val, bg=sbg, bold=True, color=stxt, align=center())
        elif j == 1:
            c2 = cell(ws1, i, j, val, bg=bg, bold=True, color="1E3A8A" if "1." in row[0] else "333333")
        else:
            c2 = cell(ws1, i, j, val, bg=bg if j <= 2 else ("F8FAFC" if i % 2 == 0 else None))
    ws1.row_dimensions[i].height = 72

widths = [24,28,34,34,40,14,10,12]
for i, w in enumerate(widths, 1):
    ws1.column_dimensions[get_column_letter(i)].width = w

# ── Sheet 2: Test Plan ────────────────────────────────────────────────────────
ws2 = wb.create_sheet("Test_Plan")
ws2.sheet_view.showGridLines = False

ws2.merge_cells("A1:D1")
c = ws2["A1"]
c.value = "TEST PLAN — Task Manager v2.0"
c.font = font(bold=True, color="FFFFFF", size=14)
c.fill = fill("0F172A"); c.alignment = center()

def section(ws, row, title):
    ws.merge_cells(f"A{row}:D{row}")
    c = ws.cell(row=row, column=1, value=title)
    c.font = font(bold=True, color="FFFFFF", size=12)
    c.fill = fill("1E3A5F"); c.alignment = left(); c.border = border()
    ws.row_dimensions[row].height = 28

def plan_row(ws, row, label, value, bg=None):
    cell(ws, row, 1, label, bg=bg or "F1F5F9", bold=True, color="1E3A8A")
    ws.merge_cells(f"B{row}:D{row}")
    c = ws.cell(row=row, column=2, value=value)
    c.font = font(); c.alignment = left(); c.border = border()
    if bg: c.fill = fill(bg)
    ws.row_dimensions[row].height = 24

section(ws2, 2, "1. PROJECT INFO")
plan_row(ws2, 3, "Project", "Task Manager v2.0 — Full-stack CRUD App")
plan_row(ws2, 4, "Version", "2.0.0")
plan_row(ws2, 5, "Prepared By", "QA Team")
plan_row(ws2, 6, "Date", "2025-06-01")
plan_row(ws2, 7, "Tools", "Postman (API), k6 (Load), GitHub Actions (CI), pytest (Unit)")

section(ws2, 9, "2. SCOPE — IN")
in_scope = [
    ("Auth Module","Login / logout, token validation, 401 on bad creds"),
    ("Task CRUD","Create, Read, Update, Delete via REST API"),
    ("UI Interactions","Create/Edit modal, Delete confirm, search, filters"),
    ("Stats Dashboard","Summary counts by status; live update on changes"),
    ("Performance","Stats endpoint P95 < 200ms under 50 VUs"),
]
for i, (label, val) in enumerate(in_scope, 10):
    plan_row(ws2, i, label, val, bg="F0FDF4")

section(ws2, 16, "3. SCOPE — OUT")
out_scope = [
    ("Email notifications","Not in v2.0 scope"),
    ("Multi-user roles","Only admin user in current build"),
    ("Persistent DB","In-memory store; no PostgreSQL in this version"),
]
for i, (label, val) in enumerate(out_scope, 17):
    plan_row(ws2, i, label, val, bg="FFF7ED")

section(ws2, 21, "4. TEST TYPES")
types = [
    ("Functional","Verify all CRUD operations and auth flow","TC-001–TC-014"),
    ("Negative","Invalid inputs, bad IDs, wrong credentials","TC-002, TC-004, TC-009"),
    ("UI/UX","Confirm dialogs, error banners, filter pills","TC-010, TC-011, TC-013"),
    ("Integration","Frontend ↔ Backend stats update flow","TC-014"),
    ("Performance","k6 load test — 50 VUs for 30s","TC-015"),
]
header_row(ws2, 22, ["Type","Description","Linked TCs"])
for i, (t, d, tc) in enumerate(types, 23):
    bg = "F8FAFC" if i % 2 == 0 else None
    cell(ws2, i, 1, t, bg=bg, bold=True, color="1E3A8A")
    cell(ws2, i, 2, d, bg=bg)
    cell(ws2, i, 3, tc, bg=bg, align=center())
    ws2.merge_cells(f"B{i}:C{i}") if False else None  # no merge here

section(ws2, 29, "5. RISKS & MITIGATIONS")
risks = [
    ("CORS errors on Windows","High","Test with 127.0.0.1 instead of localhost"),
    ("Route order bug in FastAPI","Critical","Stats route defined before /{task_id}"),
    ("Missing Node.js","Medium","Document Node 18+ prerequisite in README"),
    ("In-memory data loss on restart","Low","Document limitation; seed data auto-loads"),
]
header_row(ws2, 30, ["Risk","Severity","Mitigation"])
for i, (r, s, m) in enumerate(risks, 31):
    SEV = {"Critical":"FEE2E2","High":"FFF7ED","Medium":"FEFCE8","Low":"F0FDF4"}
    bg = SEV.get(s, "F8FAFC")
    cell(ws2, i, 1, r, bg=bg)
    cell(ws2, i, 2, s, bg=bg, bold=True, color={"Critical":"991B1B","High":"9A3412","Medium":"92400E","Low":"065F46"}.get(s,"333333"), align=center())
    cell(ws2, i, 3, m, bg=bg)

for i, w in enumerate([22, 32, 28], 1):
    ws2.column_dimensions[get_column_letter(i)].width = w
ws2.column_dimensions["D"].width = 4

# ── Sheet 3: Test Execution Log ───────────────────────────────────────────────
ws3 = wb.create_sheet("Test_Execution_Log")
ws3.sheet_view.showGridLines = False
ws3.freeze_panes = "A3"

ws3.merge_cells("A1:J1")
c = ws3["A1"]
c.value = "TEST EXECUTION LOG — Task Manager v2.0"
c.font = font(bold=True, color="FFFFFF", size=14)
c.fill = fill("0F172A"); c.alignment = center()

header_row(ws3, 2, ["TC-ID","Title","Module","Priority","Tester","Exec Date","Build","Status","Defect ID","Notes"])

exec_rows = [
    ("TC-001","Valid admin login","Auth","Critical","Alice","2025-06-01","v2.0.0","Pass","—","Token returned correctly"),
    ("TC-002","Invalid login rejected","Auth","Critical","Alice","2025-06-01","v2.0.0","Pass","—","401 returned as expected"),
    ("TC-003","Create task via POST","Task Management","High","Bob","2025-06-01","v2.0.0","Pass","—","201 + ID in response"),
    ("TC-004","Create task missing title","Task Management","High","Bob","2025-06-01","v2.0.0","Pass","—","422 validation error shown"),
    ("TC-005","List all tasks","Task Management","High","Carol","2025-06-01","v2.0.0","Pass","—","All 8 seed tasks returned"),
    ("TC-006","Filter by status=Open","Task Management","Medium","Carol","2025-06-01","v2.0.0","Pass","—","Only Open tasks returned"),
    ("TC-007","Filter by priority=Critical","Task Management","Medium","Alice","2025-06-01","v2.0.0","Pass","—","Correct subset returned"),
    ("TC-008","Update task status","Task Management","High","Bob","2025-06-02","v2.0.0","Pass","—","Status updated to Closed"),
    ("TC-009","Update non-existent task","Task Management","Medium","Bob","2025-06-02","v2.0.0","Pass","—","404 with meaningful message"),
    ("TC-010","Delete with confirmation","UI","Critical","Alice","2025-06-02","v2.0.0","Pass","DEF-001","Confirm dialog now shown. DEF-001 fixed."),
    ("TC-011","Cancel delete unchanged","UI","High","Alice","2025-06-02","v2.0.0","Pass","—","Task preserved on cancel"),
    ("TC-012","Stats summary endpoint","Dashboard","High","Carol","2025-06-02","v2.0.0","Pass","DEF-003","Route order fixed. DEF-003 resolved."),
    ("TC-013","Search by keyword","UI","Medium","Bob","2025-06-02","v2.0.0","Pass","—","Filter works on title + description"),
    ("TC-014","Stats update after add","Dashboard","High","Carol","2025-06-03","v2.0.0","In Progress","—","Frontend re-fetches on create"),
    ("TC-015","Performance load test","Performance","Low","Alice","2025-06-03","v2.0.0","Pending","—","k6 test scheduled"),
]

STATUS_BG = {"Pass":"D1FAE5","Fail":"FEE2E2","In Progress":"F3E8FF","Pending":"F1F5F9","Blocked":"FEF3C7"}
STATUS_TXT = {"Pass":"065F46","Fail":"991B1B","In Progress":"6B21A8","Pending":"64748B","Blocked":"92400E"}
PRI_COLOR  = {"Critical":"EF4444","High":"F97316","Medium":"EAB308","Low":"22C55E"}

for i, row in enumerate(exec_rows, 3):
    row_bg = "F8FAFC" if i % 2 == 0 else None
    for j, val in enumerate(row, 1):
        if j == 8:
            c2 = cell(ws3, i, j, val, bg=STATUS_BG.get(val,"F8FAFC"), bold=True, color=STATUS_TXT.get(val,"333333"), align=center())
        elif j == 4:
            c2 = cell(ws3, i, j, val, bg=row_bg, bold=True, color=PRI_COLOR.get(val,"333333"))
        else:
            c2 = cell(ws3, i, j, val, bg=row_bg)
    ws3.row_dimensions[i].height = 40

# Summary row
r = len(exec_rows) + 3
ws3.merge_cells(f"A{r}:D{r}")
cell(ws3, r, 1, "EXECUTION SUMMARY", bg="0F172A", bold=True, color="FFFFFF", align=center())
cell(ws3, r, 5, "Pass Rate:", bg="1E293B", bold=True, color="94A3B8")
cell(ws3, r, 6, f'=TEXT(COUNTIF(H3:H{r-1},"Pass")/COUNTA(H3:H{r-1}),"0.0%")', bg="ECFDF5", bold=True, color="065F46", align=center())
cell(ws3, r, 7, "Passed:", bg="1E293B", bold=True, color="22C55E")
cell(ws3, r, 8, f'=COUNTIF(H3:H{r-1},"Pass")', bg="D1FAE5", bold=True, color="065F46", align=center())
cell(ws3, r, 9, "Pending:", bg="1E293B", bold=True, color="94A3B8")
cell(ws3, r, 10, f'=COUNTIF(H3:H{r-1},"Pending")', bg="F1F5F9", bold=True, color="64748B", align=center())

ws3.row_dimensions[2].height = 32
ws3.row_dimensions[r].height = 30

widths3 = [8,28,16,10,10,12,9,12,10,32]
for i, w in enumerate(widths3, 1):
    ws3.column_dimensions[get_column_letter(i)].width = w

out = "/home/claude/task_manager/testing/STLC/STLC_Document.xlsx"
wb.save(out)
print(f"Saved: {out}")
